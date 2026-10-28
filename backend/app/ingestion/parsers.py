from __future__ import annotations

import hashlib
from dataclasses import dataclass
from pathlib import Path

import fitz
from docx import Document as DocxDocument


SUPPORTED_DOCUMENT_EXTENSIONS = {".txt", ".md", ".markdown", ".pdf", ".docx"}


@dataclass(frozen=True)
class ParsedChunk:
    text: str
    page_or_segment: str


@dataclass(frozen=True)
class ParsedDocument:
    title: str | None
    page_count: int | None
    parser: str
    chunks: list[ParsedChunk]


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def chunk_hash(text: str, marker: str) -> str:
    return hashlib.sha256(f"{marker}\n{text}".encode("utf-8", errors="ignore")).hexdigest()


def parse_document(path: Path) -> ParsedDocument:
    suffix = path.suffix.lower()
    if suffix in {".txt", ".md", ".markdown"}:
        return _parse_text(path)
    if suffix == ".pdf":
        return _parse_pdf(path)
    if suffix == ".docx":
        return _parse_docx(path)
    raise ValueError(f"Unsupported document type: {suffix}")


def _parse_text(path: Path) -> ParsedDocument:
    text = path.read_text(encoding="utf-8", errors="ignore")
    chunks = _split_text(text, marker_prefix="seg")
    return ParsedDocument(title=path.stem, page_count=1, parser="text", chunks=chunks)


def _parse_pdf(path: Path) -> ParsedDocument:
    chunks: list[ParsedChunk] = []
    with fitz.open(path) as doc:
        title = doc.metadata.get("title") or path.stem
        for index, page in enumerate(doc, start=1):
            page_text = page.get_text("text").strip()
            for chunk_index, chunk in enumerate(_split_plain(page_text), start=1):
                chunks.append(ParsedChunk(text=chunk, page_or_segment=f"p.{index}.{chunk_index}"))
        return ParsedDocument(title=title, page_count=doc.page_count, parser="pymupdf", chunks=chunks)


def _parse_docx(path: Path) -> ParsedDocument:
    doc = DocxDocument(path)
    paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    text = "\n".join(paragraphs)
    chunks = _split_text(text, marker_prefix="seg")
    return ParsedDocument(title=path.stem, page_count=None, parser="python-docx", chunks=chunks)


def _split_text(text: str, marker_prefix: str) -> list[ParsedChunk]:
    return [
        ParsedChunk(text=chunk, page_or_segment=f"{marker_prefix}.{index}")
        for index, chunk in enumerate(_split_plain(text), start=1)
    ]


def _split_plain(text: str, max_chars: int = 1400, overlap: int = 180) -> list[str]:
    normalized = "\n".join(line.strip() for line in text.splitlines())
    normalized = "\n".join(part for part in normalized.splitlines() if part)
    if not normalized:
        return []

    chunks: list[str] = []
    start = 0
    while start < len(normalized):
        end = min(start + max_chars, len(normalized))
        if end < len(normalized):
            breakpoint = max(normalized.rfind("\n", start, end), normalized.rfind("。", start, end))
            if breakpoint > start + max_chars // 2:
                end = breakpoint + 1
        chunk = normalized[start:end].strip()
        if chunk:
            chunks.append(chunk)
        if end >= len(normalized):
            break
        start = max(0, end - overlap)
    return chunks
